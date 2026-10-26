"""Functional tests for the calculations, AI anomaly rules and reports."""
import io

from openpyxl import load_workbook
from pypdf import PdfReader

from ai import detect_anomalies, rule_based_summary
from calculations import analyse_item, analyse_project
from conftest import create_project, post_form, register
from test_security import SAMPLE_CSV, upload


def test_variance_formulas():
    assert analyse_item(100, 150) == {"variance": 50, "pct_variance": 50.0, "status": "Over budget"}
    assert analyse_item(100, 80) == {"variance": -20, "pct_variance": -20.0, "status": "Under budget"}
    assert analyse_item(100, 101.5)["status"] == "On budget"  # within +/-2%
    assert analyse_item(0, 500) == {"variance": 500, "pct_variance": None, "status": "Over budget"}


def test_project_totals_match_hand_calculation():
    import pandas as pd
    rows = pd.read_csv(io.BytesIO(SAMPLE_CSV)).rename(columns=str.lower).to_dict("records")
    t = analyse_project(rows)["totals"]
    assert (t["estimated"], t["actual"], t["variance"], t["pct_variance"]) == (2_085_000, 2_222_500, 137_500, 6.59)


def test_anomaly_detection_flags_the_obvious_overruns():
    import pandas as pd
    rows = pd.read_csv(io.BytesIO(SAMPLE_CSV)).rename(columns=str.lower).to_dict("records")
    flagged = {a["item"]: a["severity"] for a in detect_anomalies(analyse_project(rows)["lines"])}
    assert flagged == {"Network setup": "High", "Miscellaneous supplies": "High"}


def test_rule_based_summary_without_api_key():
    lines = [{"item": "A", "category": "X", "estimated": 100.0, "actual": 200.0}]
    a = analyse_project(lines)
    text = rule_based_summary("Demo", a, detect_anomalies(a["lines"]))
    assert "over budget by 100.00 (+100.0%)" in text


def test_pdf_and_excel_reports(client):
    register(client, "reporter")
    pid = create_project(client, "Report test")
    upload(client, pid, "costs.csv", SAMPLE_CSV)
    pdf = client.get(f"/projects/{pid}/report.pdf")
    assert pdf.status_code == 200 and pdf.mimetype == "application/pdf"
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    assert "Cost Report: Report test" in text and "2,222,500.00" in text
    wb = load_workbook(io.BytesIO(client.get(f"/projects/{pid}/report.xlsx").data))
    assert wb.sheetnames == ["Summary", "Items", "Categories", "Unusual expenses"]
    assert wb["Items"].max_row == 11


def test_manual_entry_and_delete(client):
    register(client, "editor")
    pid = create_project(client)
    r = post_form(client, f"/projects/{pid}/items", {"item": "Desk", "estimated": "1,000", "actual": "1200"},
                  follow_redirects=True)
    assert b"Item added" in r.data and b"+200.00" in r.data
