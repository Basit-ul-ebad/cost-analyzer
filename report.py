"""Downloadable reports: a PDF cost report and an Excel export of the results."""
import io
import re
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import matplotlib
matplotlib.use("Agg")  # no display needed on a server
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

BLUE = "#2457d6"
GREY = "#93a4c3"
RED = "#d14343"
GREEN = "#2e9d57"
MUTED = "#9aa3b2"
STATUS_COLOR = {"Over budget": RED, "Under budget": GREEN, "On budget": MUTED, "Critical": RED}


def safe_filename(name, ext):
    slug = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-")[:50] or "project"
    return f"{slug}-cost-report.{ext}"


def _money(v):
    return f"{v:,.2f}"


def _signed(v):
    return f"{v:+,.2f}"


def _pct(v):
    return "n/a" if v is None else f"{v:+.1f}%"


def _short(v, _pos=None):
    a = abs(v)
    if a >= 1e6:
        return f"{v / 1e6:.1f}M"
    if a >= 1e3:
        return f"{v / 1e3:.0f}K"
    return f"{v:.0f}"


# ------------------------------------------------------------------ charts
def _fig_to_png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _trim(label, n=22):
    return label if len(label) <= n else label[: n - 1] + "…"


def chart_category(analysis):
    cats = sorted(analysis["categories"], key=lambda c: c["actual"], reverse=True)[:10]
    labels = [_trim(c["category"], 14) for c in cats]
    x = range(len(cats))
    fig, ax = plt.subplots(figsize=(7.2, 3.2))
    w = 0.38
    ax.bar([i - w / 2 for i in x], [c["estimated"] for c in cats], w, label="Estimated", color=GREY)
    ax.bar([i + w / 2 for i in x], [c["actual"] for c in cats], w, label="Actual", color=BLUE)
    ax.set_xticks(list(x), labels, rotation=25, ha="right", fontsize=8)
    ax.yaxis.set_major_formatter(FuncFormatter(_short))
    ax.tick_params(axis="y", labelsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#e5e7eb", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=8, loc="upper right")
    ax.set_title("Estimated vs actual cost by category", fontsize=10, loc="left")
    return _fig_to_png(fig)


def chart_variance(analysis):
    lines = sorted(analysis["lines"], key=lambda l: l["variance"])[-15:]  # up to 15 largest
    if len(analysis["lines"]) > 15:
        lines = sorted(analysis["lines"], key=lambda l: abs(l["variance"]))[-15:]
        lines.sort(key=lambda l: l["variance"])
    fig, ax = plt.subplots(figsize=(7.2, max(2.4, 0.28 * len(lines) + 0.8)))
    ax.barh([_trim(l["item"]) for l in lines], [l["variance"] for l in lines],
            color=[STATUS_COLOR[l["status"]] for l in lines])
    ax.axvline(0, color="#6b7280", linewidth=0.8)
    ax.xaxis.set_major_formatter(FuncFormatter(_short))
    ax.tick_params(labelsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#e5e7eb", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.set_title("Variance per item (red = over, green = under, grey = on budget)", fontsize=10, loc="left")
    return _fig_to_png(fig)


# ------------------------------------------------------------------ PDF
def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=20, leading=24, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#6b7280")),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=13, spaceBefore=12, spaceAfter=6, keepWithNext=1,
                             textColor=colors.HexColor("#111827")),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=14),
        "small": ParagraphStyle("sm", parent=ss["Normal"], fontSize=8, leading=11,
                                textColor=colors.HexColor("#6b7280")),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=8, leading=10),
        "cellr": ParagraphStyle("cr", parent=ss["Normal"], fontSize=8, leading=10, alignment=TA_RIGHT),
    }


def _p(text, style):
    """Paragraph with user text escaped (reportlab parses <tags> in Paragraph text)."""
    return Paragraph(escape(str(text)), style)


def _table(rows, col_widths, num_cols, status_col=None, total_row=False):
    t = Table(rows, colWidths=col_widths, repeatRows=1)
    style = [
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#4b5563")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f4f9")),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#e3e6eb")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for c in num_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    for r in range(1, len(rows)):
        if r % 2 == 0:
            style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#fafbfc")))
    if total_row:
        style += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8),
                  ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#9aa3b2"))]
    t.setStyle(TableStyle(style))
    return t


def _status_para(status, st):
    color = STATUS_COLOR.get(status, MUTED)
    return Paragraph(f'<font color="{color}"><b>{escape(status)}</b></font>', st["cell"])


def build_pdf(project_name, analysis, anomalies, summary, pert=None):
    st = _styles()
    t = analysis["totals"]
    buf = io.BytesIO()
    generated = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#9aa3b2"))
        canvas.drawString(18 * mm, 10 * mm, f"Cost Analyzer · {project_name[:60]}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=18 * mm,
                            title=f"Cost report - {project_name}", author="Cost Analyzer")
    W = A4[0] - 36 * mm
    story = [
        _p(f"Cost Report: {project_name}", st["title"]),
        _p(f"Generated {generated} · {t['item_count']} cost items", st["sub"]),
        Spacer(1, 10),
    ]

    # Summary cards
    color = STATUS_COLOR.get(t["status"], MUTED)
    cards = [[
        Paragraph(f'<font size="7.5" color="#6b7280">ESTIMATED COST</font><br/><font size="13"><b>{_money(t["estimated"])}</b></font>', st["body"]),
        Paragraph(f'<font size="7.5" color="#6b7280">ACTUAL COST</font><br/><font size="13"><b>{_money(t["actual"])}</b></font>', st["body"]),
        Paragraph(f'<font size="7.5" color="#6b7280">VARIANCE</font><br/><font size="13" color="{color}"><b>{_signed(t["variance"])}</b></font>', st["body"]),
        Paragraph(f'<font size="7.5" color="#6b7280">% VARIANCE</font><br/><font size="13" color="{color}"><b>{_pct(t["pct_variance"])}</b></font> <font size="8" color="{color}">{t["status"]}</font>', st["body"]),
    ]]
    ct = Table(cards, colWidths=[W / 4] * 4, rowHeights=[38])
    ct.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#e3e6eb")),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#e3e6eb")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    sc = t["status_counts"]
    story += [ct, Spacer(1, 4), _p(
        f"{sc['Over budget']} over budget · {sc['On budget']} on budget · {sc['Under budget']} under budget "
        f"(on budget = within ±{analysis['tolerance_pct']:.0f}% of estimate). "
        "Variance = Actual − Estimated; % Variance = Variance ÷ Estimated × 100.", st["small"])]

    # AI summary
    story.append(_p("Cost analysis summary", st["h2"]))
    if summary:
        for para in summary["summary"].split("\n"):
            story.append(_p(para, st["body"]) if para.strip() else Spacer(1, 5))
        story.append(Spacer(1, 4))
        story.append(_p(f"Source: {summary['source']} · generated {summary['generated_at']} UTC. "
                        "AI-generated text may contain mistakes; verify against the tables.", st["small"]))
        if summary.get("stale"):
            story.append(_p("Note: cost data changed after this summary was generated.", st["small"]))
    else:
        story.append(_p("No summary has been generated for this project yet.", st["small"]))

    # Unusual expenses
    story.append(_p("Unusual expenses", st["h2"]))
    if anomalies:
        rows = [["Item", "Category", "Variance", "% Var.", "Severity", "Why flagged"]]
        for a in anomalies:
            rows.append([_p(a["item"], st["cell"]), _p(a["category"], st["cell"]),
                         _signed(a["variance"]), _pct(a["pct_variance"]), a["severity"],
                         _p("; ".join(a["reasons"]), st["cell"])])
        story.append(_table(rows, [W * .19, W * .13, W * .13, W * .1, W * .09, W * .36], [2, 3]))
    else:
        story.append(_p("No unusual expenses detected.", st["body"]))

    # Charts
    story += [PageBreak(), _p("Charts", st["h2"])]
    story.append(Image(chart_category(analysis), width=W, height=W * 3.2 / 7.2 * 1.08, kind="proportional"))
    story.append(Spacer(1, 8))
    vimg = chart_variance(analysis)
    n = min(len(analysis["lines"]), 15)
    vh = max(2.4, 0.28 * n + 0.8)
    story.append(Image(vimg, width=W, height=W * vh / 7.2 * 1.08, kind="proportional"))
    if len(analysis["lines"]) > 15:
        story.append(_p("Variance chart shows the 15 items with the largest variance.", st["small"]))

    # Category table
    story.append(_p("Costs by category", st["h2"]))
    rows = [["Category", "Items", "Estimated", "Actual", "Variance", "% Var.", "Status"]]
    for c in analysis["categories"]:
        rows.append([_p(c["category"], st["cell"]), c["count"], _money(c["estimated"]), _money(c["actual"]),
                     _signed(c["variance"]), _pct(c["pct_variance"]), _status_para(c["status"], st)])
    story.append(_table(rows, [W * .22, W * .07, W * .15, W * .15, W * .14, W * .1, W * .17], [1, 2, 3, 4, 5]))

    # Item table
    story.append(_p("All cost items", st["h2"]))
    rows = [["Item", "Category", "Estimated", "Actual", "Variance", "% Var.", "Status"]]
    for l in analysis["lines"]:
        rows.append([_p(l["item"], st["cell"]), _p(l["category"], st["cell"]), _money(l["estimated"]),
                     _money(l["actual"]), _signed(l["variance"]), _pct(l["pct_variance"]),
                     _status_para(l["status"], st)])
    rows.append(["Total", "", _money(t["estimated"]), _money(t["actual"]), _signed(t["variance"]),
                 _pct(t["pct_variance"]), _status_para(t["status"], st)])
    story.append(_table(rows, [W * .22, W * .13, W * .14, W * .14, W * .14, W * .1, W * .13],
                        [2, 3, 4, 5], total_row=True))

    story += _pert_story(pert or {}, st, W)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def chart_gantt(sched):
    acts = sched["activities"][:30]
    fig, ax = plt.subplots(figsize=(7.2, max(2.2, 0.26 * len(acts) + 0.9)))
    labels = [f"{a['code']}  {_trim(a['item'], 26)}" for a in acts]
    ax.barh(labels, [a["ef"] - a["es"] for a in acts], left=[a["es"] for a in acts],
            color=[RED if a["critical"] else GREY for a in acts], height=0.6)
    ax.invert_yaxis()
    ax.set_xlabel("Days from start", fontsize=8)
    ax.tick_params(labelsize=7.5)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#e5e7eb", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.set_title("PERT schedule (red = critical path)", fontsize=10, loc="left")
    return _fig_to_png(fig)


def _pert_story(pert, st, W):
    cost, sched = pert.get("cost"), pert.get("schedule")
    if not cost and not sched:
        return []
    out = [PageBreak(), _p("PERT analysis", st["h2"]),
           _p("Three-point estimates: Expected E = (O + 4M + P) / 6, standard deviation = (P - O) / 6. "
              "Probabilities assume a normal distribution of the total.", st["small"])]
    if cost:
        out.append(_p("Cost PERT", st["h2"]))
        rows = [["Measure", "Value"],
                ["Items with your own three-point costs", str(cost["given"])],
                ["Items with auto-estimated ranges", str(cost["auto"])],
                ["Expected total cost (E)", _money(cost["expected"])],
                ["Standard deviation", _money(cost["sd"])],
                ["Budget (sum of estimates)", _money(cost["budget"])],
                ["Probability of staying within budget", f"{cost['p_within_budget']}%"],
                ["Actual cost so far", _money(cost["actual"])]]
        rows += [[f"Cost at {c['level']}% confidence", _money(c["value"])] for c in cost["confidence"]]
        out.append(_table(rows, [W * .55, W * .45], [1]))
        if cost["auto"]:
            out.append(_p(f"Auto-estimated ranges: most likely = the estimate; optimistic = estimate x "
                          f"{cost['auto_factors']['o']:.2f}, pessimistic = estimate x {cost['auto_factors']['p']:.2f}, "
                          f"based on {cost['auto_basis']}. Rows marked * use these ranges.", st["small"]))
        out.append(Spacer(1, 6))
        rows = [["ID", "Item", "O", "M", "P", "Expected", "SD", "Actual"]]
        for i in cost["items"][:25]:
            rows.append([i["code"], _p(i["item"] + (" *" if i["auto"] else ""), st["cell"]), _money(i["o"]), _money(i["m"]), _money(i["p"]),
                         _money(i["expected"]), _money(i["sd"]), _money(i["actual"])])
        out.append(_table(rows, [W * .07, W * .25, W * .11, W * .11, W * .11, W * .12, W * .1, W * .13],
                          [2, 3, 4, 5, 6, 7]))
        if len(cost["items"]) > 25:
            out.append(_p("Showing the 25 items with the largest uncertainty.", st["small"]))
    if sched:
        out.append(_p("Schedule PERT", st["h2"]))
        if sched.get("error"):
            out.append(_p(sched["error"], st["body"]))
            return out
        rows = [["Measure", "Value"],
                ["Expected project duration (T)", f"{sched['duration']:.1f} days"],
                ["Standard deviation", f"{sched['sd']:.1f} days"],
                ["Critical path", " -> ".join(sched["critical_path"])]]
        if sched.get("target_days"):
            rows.append([f"Probability of finishing within {sched['target_days']:g} days",
                         f"{sched['p_target']}% (Z = {sched['z_target']})"])
        rows += [[f"Duration at {c['level']}% confidence", f"{c['value']} days"] for c in sched["confidence"]]
        out.append(_table([[_p(str(a), st["cell"]) if i == 1 else a for i, a in enumerate(r)] for r in rows],
                          [W * .55, W * .45], []))
        out.append(Spacer(1, 6))
        n = min(len(sched["activities"]), 30)
        out.append(Image(chart_gantt(sched), width=W, height=W * max(2.2, 0.26 * n + 0.9) / 7.2 * 1.08,
                         kind="proportional"))
        rows = [["ID", "Activity", "Pred.", "te", "SD", "ES", "EF", "LS", "LF", "Slack", ""]]
        for a in sched["activities"]:
            rows.append([a["code"], _p(a["item"], st["cell"]), _p(", ".join(a["preds"]) or "-", st["cell"]),
                         f"{a['te']:.2f}", f"{a['sd']:.2f}", f"{a['es']:.1f}", f"{a['ef']:.1f}",
                         f"{a['ls']:.1f}", f"{a['lf']:.1f}", f"{a['slack']:.1f}",
                         _status_para("Critical", st) if a["critical"] else ""])
        out.append(_table(rows, [W * .06, W * .25, W * .1, W * .07, W * .07, W * .07, W * .07, W * .07,
                                 W * .07, W * .07, W * .1], [3, 4, 5, 6, 7, 8, 9]))
        for w in sched.get("warnings", [])[:10]:
            out.append(_p("Note: " + w, st["small"]))
    return out


# ------------------------------------------------------------------ Excel
def _xl_safe(value):
    """Stop spreadsheet formula injection: text starting with = + - @ is shown as text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def build_xlsx(project_name, analysis, anomalies, summary, pert=None):
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2457D6")

    def sheet(ws, headers, rows, widths, money_cols=()):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in rows:
            ws.append([_xl_safe(v) for v in r])
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0.00"
        ws.freeze_panes = "A2"

    t = analysis["totals"]
    ws = wb.active
    ws.title = "Summary"
    ws.append([_xl_safe(f"Cost report: {project_name}")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {datetime.now(timezone.utc):%d %b %Y %H:%M} UTC"])
    ws.append([])
    for label, val in [("Total estimated", t["estimated"]), ("Total actual", t["actual"]),
                       ("Variance", t["variance"]), ("% Variance", t["pct_variance"]),
                       ("Status", t["status"]), ("Items", t["item_count"])]:
        ws.append([label, val])
    for r in range(4, 7):
        ws.cell(r, 2).number_format = "#,##0.00"
    ws.append([])
    ws.append(["AI summary"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    if summary:
        ws.append([_xl_safe(summary["summary"])])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 300
        ws.append([f"Source: {summary['source']}"])
    else:
        ws.append(["Not generated yet"])
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 18

    sheet(wb.create_sheet("Items"),
          ["Item", "Category", "Estimated", "Actual", "Variance", "% Variance", "Status"],
          [[l["item"], l["category"], l["estimated"], l["actual"], l["variance"], l["pct_variance"], l["status"]]
           for l in analysis["lines"]],
          [32, 18, 15, 15, 15, 12, 14], money_cols=(3, 4, 5))
    sheet(wb.create_sheet("Categories"),
          ["Category", "Items", "Estimated", "Actual", "Variance", "% Variance", "Status"],
          [[c["category"], c["count"], c["estimated"], c["actual"], c["variance"], c["pct_variance"], c["status"]]
           for c in analysis["categories"]],
          [22, 8, 15, 15, 15, 12, 14], money_cols=(3, 4, 5))
    sheet(wb.create_sheet("Unusual expenses"),
          ["Item", "Category", "Variance", "% Variance", "Severity", "Reasons"],
          [[a["item"], a["category"], a["variance"], a["pct_variance"], a["severity"], "; ".join(a["reasons"])]
           for a in anomalies],
          [32, 18, 15, 12, 10, 80], money_cols=(3,))
    cost, sched = (pert or {}).get("cost"), (pert or {}).get("schedule")
    if cost:
        ws = wb.create_sheet("PERT cost")
        sheet(ws, ["ID", "Item", "Optimistic", "Most likely", "Pessimistic", "Expected", "Std dev", "Actual", "Range source"],
              [[i["code"], i["item"], i["o"], i["m"], i["p"], i["expected"], i["sd"], i["actual"],
                "auto" if i["auto"] else "entered"] for i in cost["items"]],
              [8, 32, 14, 14, 14, 14, 12, 14, 12], money_cols=(3, 4, 5, 6, 7, 8))
        ws.append([])
        for label, val in [("Expected total cost", cost["expected"]), ("Standard deviation", cost["sd"]),
                           ("Budget", cost["budget"]), ("P(within budget) %", cost["p_within_budget"]),
                           ("Auto-range basis", cost["auto_basis"] or "none (all ranges entered)")] + \
                          [(f"Cost at {c['level']}% confidence", c["value"]) for c in cost["confidence"]]:
            ws.append([None, label, val])
    if sched and not sched.get("error"):
        ws = wb.create_sheet("PERT schedule")
        sheet(ws, ["ID", "Activity", "Predecessors", "O", "M", "P", "te", "SD", "ES", "EF", "LS", "LF", "Slack", "Critical"],
              [[a["code"], a["item"], ", ".join(a["preds"]), a["o"], a["m"], a["p"], a["te"], a["sd"], a["es"], a["ef"],
                a["ls"], a["lf"], a["slack"], "Yes" if a["critical"] else ""] for a in sched["activities"]],
              [8, 32, 14, 7, 7, 7, 8, 8, 8, 8, 8, 8, 8, 9])
        ws.append([])
        rows = [("Expected duration (days)", sched["duration"]), ("Standard deviation (days)", sched["sd"]),
                ("Critical path", " -> ".join(sched["critical_path"]))]
        if sched.get("target_days"):
            rows += [("Target (days)", sched["target_days"]), ("P(finish by target) %", sched["p_target"])]
        for label, val in rows:
            ws.append([None, label, _xl_safe(val)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
